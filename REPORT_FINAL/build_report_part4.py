"""Part 4: rework report for the Streamlit application."""
import docx, os
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"C:\Users\Prajvesh\Documents\Database-Management-System-for-E-Commerce-Order-Return-and-Refund-Management\REPORT_FINAL\25WU0102114_Prajvesh_DBMS_PBL_Report.docx"
SS = r"C:\Users\Prajvesh\Documents\Database-Management-System-for-E-Commerce-Order-Return-and-Refund-Management\ecommerce_refund_dbms\docs\screenshots"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
d = docx.Document(OUT)


def has_image(p):
    return len(p._element.findall(".//%sdrawing" % W)) > 0


def replace_image_before_caption(caption_start, image_path, width=6.0):
    for j, p in enumerate(d.paragraphs):
        if p.text.strip().startswith(caption_start):
            for k in range(j - 1, max(j - 4, -1), -1):
                if has_image(d.paragraphs[k]):
                    d.paragraphs[k].text = ""
                    r = d.paragraphs[k].add_run()
                    r.add_picture(image_path, width=Inches(width))
                    d.paragraphs[k].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    return True
    return False


def set_caption(caption_start, new_text):
    for p in d.paragraphs:
        if p.text.strip().startswith(caption_start):
            p.text = new_text
            return True
    return False


def add_row(ti, vals):
    t = d.tables[ti]
    row = t.add_row()
    for i, v in enumerate(vals):
        if i < len(row.cells):
            row.cells[i].text = v

# ---- Abstract (keep 150-200 words) ----
abstract = ("E-commerce post-purchase operations routinely fail when orders, returns, and refunds are tracked in disconnected spreadsheets, "
"causing untracked status changes, duplicate refunds, and missing audit evidence. The objective of this project is to design and implement a centralized ACID-compliant DBMS for the complete order-to-refund lifecycle. "
"The solution uses a 13-table Third Normal Form (3NF) MySQL 8 (InnoDB) schema covering customers, products, inventory, orders, payments, item-level returns, warehouse inspection, refunds, and an immutable audit log; five triggers enforce return-window eligibility, return-quantity limits, strict state transitions, and automatic auditing; and one stored procedure, sp_process_refund, executes atomic refund processing with duplicate-refund prevention and full rollback on error [5], [7]. "
"A role-based Streamlit application connects directly to the database with parameterized SQL to implement ordering, returning, approving, receiving, inspecting, refunding, auditing, and analytics, while the FastAPI service is retained as the automated test and benchmark harness [6], [9]. "
"The completed system demonstrates the full lifecycle Order \u2192 Requested \u2192 Approved \u2192 Received \u2192 Inspected \u2192 Refunded, passes 13/13 automated tests, sustains sub-millisecond indexed lookups at 1K scale, and verifies backup/restore integrity at 111,025 orders. "
"The application provides a single queryable source of truth for post-purchase operations with enforced business rules and traceable refunds.")
print("abstract words:", len(abstract.split()))
for p in d.paragraphs:
    if p.text.startswith("E-commerce post-purchase operations routinely fail"):
        p.text = abstract
        break
for p in d.paragraphs:
    if p.text.startswith("Keywords:"):
        p.text = "Keywords: E-Commerce Returns, Streamlit, MySQL, 3NF Normalization, Stored Procedures, Triggers"
        break

# ---- Objective 3 ----
for p in d.paragraphs:
    if p.text.startswith("Develop a FastAPI"):
        p.text = ("Develop a role-based Streamlit application with direct parameterized-SQL connectivity and session-state RBAC "
"for Customer, Support, Warehouse, and Admin, covering ordering, returns, inspection, refunds, analytics, and audit, while retaining FastAPI as the automated test/benchmark harness.")
        break

# ---- Table 11 modules row ----
t11 = d.tables[11]
for r in t11.rows:
    if r.cells[0].text.startswith("Major Application Modules"):
        r.cells[1].text = ("Streamlit pages: sidebar login + session-state RBAC; Customer Portal (catalogue, orders, returns); "
"Operations (support approve/reject, warehouse receive/inspect); Refunds (sp_process_refund); Analytics (metrics, charts, 6 reports); "
"Audit viewer (Admin); 7-state Lifecycle Demo stepper.")
    if r.cells[0].text.startswith("Proposed Workflow"):
        r.cells[1].text = ("Customer orders (Delivered) \u2192 requests return (Requested) \u2192 Support approves (Approved) \u2192 Warehouse receives "
"(Received) \u2192 inspects (Inspected) \u2192 Support/Admin refunds (Refunded) via Streamlit pages; Rejected is terminal; triggers block any skip.")

# ---- 3.1 overview ----
for p in d.paragraphs:
    if p.text.startswith("The application is a three-layer system"):
        p.text = ("The application was re-implemented in Streamlit as a role-based multi-page app that connects directly to MySQL 8 "
"with parameterized SQL in streamlit_app/db_ops.py: login is verified with bcrypt against AppUser\u22c8Role and the session role gates six pages "
"(Customer Portal, Operations, Refunds, Analytics, Audit Log, Lifecycle Demo). Every screen operation maps to the 3NF schema in Chapter 2: ordering inserts "
"Orders/OrderItem/Payment in one explicit transaction, returning inserts ReturnRequest/ReturnItem (trigger-validated), approvals/receives/inspections are constrained "
"UPDATEs/INSERTs, and refunds execute only through sp_process_refund. Analytics screens run the six catalogued JOIN/GROUP BY queries with st.metric/st.bar_chart/st.dataframe, "
"and the Admin audit page reads AuditLog. The FastAPI service is retained unchanged as the automated test and benchmark harness (17 endpoints), so all API evidence stays valid [6], [9]. "
"No screen uses isolated or static data.")
        break

# ---- Table 19 env ----
t19 = d.tables[19]
t19.rows[3].cells[1].text = "Streamlit 1.65 (application) + FastAPI/Uvicorn (retained test & benchmark harness); bcrypt; session-state RBAC"
t19.rows[4].cells[1].text = "mysql-connector-python (parameterized %s queries, explicit transactions)"
t19.rows[6].cells[1].text = ("ecommerce_refund_dbms/streamlit_app/ (app.py, db_ops.py) + backend/ (harness), "
"frontend/index.html (legacy SPA), tests/test_streamlit_ops.py")

# ---- Table 20 module mapping -> Streamlit pages ----
mods = [
    ("Sidebar login + RBAC", "bcrypt verify; session role gates pages", "Role, AppUser", "SELECT AppUser\u22c8Role; st.session_state routing; wrong password rejected"),
    ("Customer Portal", "Browse, buy, track, return", "Product, Orders, OrderItem, Payment, ReturnRequest, ReturnItem", "st.form order/return; st.dataframe tracker; explicit COMMIT txn"),
    ("Operations", "Approve/reject; receive; inspect", "ReturnRequest, ReturnItem, Inspection", "st buttons + inspect st.form; trigger blocks illegal jumps"),
    ("Refunds", "Exactly-once refund via procedure", "ReturnRequest, Refund, Payment", "st button calls sp_process_refund; duplicate blocked"),
    ("Analytics", "Metrics, charts, 6 reports (a-f)", "All sales/return tables", "st.metric + st.bar_chart + st.dataframe"),
    ("Audit Log (Admin)", "Immutable trail viewer", "AuditLog", "st.dataframe last 100; page hidden for non-admin roles"),
    ("Lifecycle Demo stepper", "Guided 7-state end-to-end run", "Full chain incl. Inspection/Refund/AuditLog", "st.status steps + graphviz state map"),
]
t20 = d.tables[20]
for i in range(1, len(t20.rows)):
    for c in t20.rows[i].cells:
        c.text = ""
for i, vals in enumerate(mods):
    if i + 1 < len(t20.rows):
        for j, v in enumerate(vals):
            t20.rows[i + 1].cells[j].text = v
    else:
        add_row(20, list(vals))

# ---- Table 21 CRUD evidence ----
t21 = d.tables[21]
t21.rows[1].cells[1].text = "Customer Portal st.forms: Place Order; Submit Return"
t21.rows[1].cells[3].text = "Fig. 2 Streamlit portal; db_ops.place_order/request_return (explicit COMMIT)"
t21.rows[2].cells[3].text = "st.dataframe tracker/queue/analytics + st.bar_chart; Figs. 2-4"
t21.rows[3].cells[3].text = "st buttons + inspect form (Fig. 3); trigger blocks skips; lifecycle demo"
t21.rows[4].cells[3].text = "Session-role gating (pages hidden \u2248 403); audit viewer Fig. 5; TC-05 RESTRICT"

# ---- Table 23 report outputs ----
outs = ["Full table in Customer Portal tracker", "Worklist in Operations queue",
        "Reconciliation table + Analytics view", "Bar chart in Analytics (highest first)",
        "Bar chart in Analytics (month-wise value)", "Ranked table in Analytics"]
for i, o in enumerate(outs):
    t23 = d.tables[23]
    t23.rows[i + 1].cells[2].text = o

# ---- Figures 2-5: swap to Streamlit screenshots ----
shots = [
    ("Figure 2.", os.path.join(SS, "st_customer_portal.png"),
     "Figure 2. Streamlit Customer Portal \u2014 catalogue, place-order form, status tracker and return request (customer_user)"),
    ("Figure 3.", os.path.join(SS, "st_ops_dashboard.png"),
     "Figure 3. Streamlit Operations Dashboard \u2014 approve/reject, receive and inspection with live queue (admin_user)"),
    ("Figure 4.", os.path.join(SS, "st_analytics.png"),
     "Figure 4. Streamlit Analytics Dashboard \u2014 refund metrics, return-rate and monthly-refund charts, top reasons"),
    ("Figure 5.", os.path.join(SS, "st_audit_log.png"),
     "Figure 5. Streamlit Admin Audit Log \u2014 immutable trigger-written trail viewer"),
    ("Figure 8.", os.path.join(SS, "test_run_output.png"),
     "Figure 8. Automated test evidence \u2014 pytest 13/13 passed (9 API + 4 Streamlit ops tests)"),
]
for cap, img, newcap in shots:
    ok = replace_image_before_caption(cap, img)
    set_caption(cap, newcap)
    print(cap, "image replaced:", ok)

# ---- Table 24: add TC-10..TC-13 ----
new_tests = [
    ("TC-10", "Streamlit auth + RBAC", "Login 4 demo roles; wrong password", "Roles resolved; invalid rejected", "Roles resolved; rejected", "Pass"),
    ("TC-11", "Streamlit lifecycle (db_ops)", "db_ops chain order\u2192refund", "Refunded + audit entries", "Refunded + audited", "Pass"),
    ("TC-12", "Streamlit guards", "Late return (order 3); 2nd refund", "Window-expiry + duplicate errors", "Both blocked as expected", "Pass"),
    ("TC-13", "Catalogue + analytics sanity", "list_products; return-rates query", ">=5 products/rates rows", "Rows returned", "Pass"),
]
for vals in new_tests:
    add_row(24, list(vals))

for p in d.paragraphs:
    if p.text.startswith("Key results: 9/9"):
        p.text = ("Key results: 13/13 tests passed (100%: 9 API + 4 Streamlit ops). Constraints (TC-01\u2013TC-05) prove the engine rejects bad quantities, "
"negative amounts, duplicate gateway references, illegal ENUM states, and restricted deletes. TC-06/TC-08 prove the procedure guards state and rolls back fully "
"on error with no orphan rows. TC-07 proves the API-to-database chain completes Order\u2192Refunded through all six states. TC-09 proves exactly-once refund under "
"concurrency (one 200, one 400). TC-10 proves Streamlit bcrypt login resolves all four RBAC roles and rejects bad passwords. TC-11 proves the Streamlit data layer "
"(db_ops) drives the same full lifecycle with explicit transactions and audit entries. TC-12 proves the late-return window trigger and duplicate-refund guard fire "
"through db_ops identically. Benchmarks (benchmarks/run_benchmarks.py): pending-returns median 0.22 ms (1K indexed) to 2.95 ms (100K indexed); sp_process_refund median "
"1.23 ms; mysqldump backup/restore of 111,025 orders verified with checksum PASS. Together this confirms both the Streamlit application and the API harness correctly "
"use the designed database and enforce all critical business rules [7].")
        break

# ---- global 9 -> 13 test-count phrases ----
for p in d.paragraphs:
    if "9/9" in p.text and not p.text.startswith("Figure 8"):
        p.text = p.text.replace("9/9", "13/13")
    if "9 automated tests" in p.text:
        p.text = p.text.replace("9 automated tests", "13 automated tests")
    if "9-test" in p.text:
        p.text = p.text.replace("9-test", "13-test")

# ---- References: add [9] ----
for p in d.paragraphs:
    if p.text.startswith("[8]"):
        from copy import deepcopy
        nel = deepcopy(p._element)
        p._element.addnext(nel)
        break
found8 = False
for p in d.paragraphs:
    if p.text.startswith("[8]"):
        if not found8:
            found8 = True
        else:
            p.text = ("[9] Streamlit Inc., Streamlit Documentation \u2014 Build data apps in Python. "
"Accessed: Oct. 11, 2026. [Online]. Available: https://docs.streamlit.io/")
            break

# ---- Appendix B ----
for p in d.paragraphs:
    if p.text.startswith("Appendix B (ecommerce_refund_dbms/"):
        p.text = ("Appendix B (ecommerce_refund_dbms/): streamlit_app/app.py (6 role-gated pages: Customer Portal, Operations, Refunds, Analytics, "
"Audit Log, Lifecycle Demo), streamlit_app/db_ops.py (parameterized operations + AppError surfacing of trigger/procedure messages), "
"backend/ retained as test/benchmark harness (17 REST endpoints: main.py, auth.py bcrypt+JWT, database.py connector, payment_gateway.py mock), "
"frontend/index.html (legacy SPA), tests/test_streamlit_ops.py (TC-10\u2013TC-13), benchmarks/run_benchmarks.py, docker-compose.yml. "
"Run: streamlit run streamlit_app/app.py --server.port 8501 \u2192 http://127.0.0.1:8501. Demo logins (password123): customer_user / support_agent / warehouse_op / admin_user.")
        break
for p in d.paragraphs:
    if p.text.startswith("Appendix C:"):
        p.text = p.text.replace("test_run_output 9/9", "test_run_output 13/13")
        if "st_customer_portal" not in p.text:
            p.text = p.text + " Streamlit screens: st_customer_portal, st_ops_dashboard, st_analytics, st_audit_log, st_lifecycle."
        break

# ---- LOF titles ----
t7 = d.tables[7]
lof = {1: "Streamlit Customer Portal", 2: "Streamlit Operations Dashboard",
       3: "Streamlit Analytics Dashboard", 4: "Streamlit Admin Audit Log"}
for idx, title in lof.items():
    if idx < len(t7.rows):
        t7.rows[idx].cells[1].text = title
# Fig 8 LOF row
for r in t7.rows:
    if r.cells[0].text.strip() == "Figure 8":
        r.cells[1].text = "Automated test evidence (13/13 passed)"

print("part4 done")
d.save(OUT)
