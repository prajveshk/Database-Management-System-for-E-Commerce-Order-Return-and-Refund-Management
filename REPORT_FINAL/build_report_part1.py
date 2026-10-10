"""Build final DBMS PBL report strictly from template."""
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy
import os

TEMPLATE = r"C:\Users\Prajvesh\Downloads\DBMS_PBL_Report_Template_Final_with_Guidelines.docx"
OUT = r"C:\Users\Prajvesh\Documents\Database-Management-System-for-E-Commerce-Order-Return-and-Refund-Management\REPORT_FINAL\25WU0102114_Prajvesh_DBMS_PBL_Report.docx"
ER_IMG = r"C:\Users\Prajvesh\Documents\Database-Management-System-for-E-Commerce-Order-Return-and-Refund-Management\ecommerce_refund_dbms\docs\er_diagram.png"
SS_DIR = r"C:\Users\Prajvesh\Documents\Database-Management-System-for-E-Commerce-Order-Return-and-Refund-Management\ecommerce_refund_dbms\docs\screenshots"
DIAG_DIR = r"C:\Users\Prajvesh\Documents\Database-Management-System-for-E-Commerce-Order-Return-and-Refund-Management\REPORT_FINAL\diagrams"

d = docx.Document(TEMPLATE)

def set_para(idx_or_text, new_text, bold=False):
    """Replace paragraph by index or by matching substring. Returns True if done."""
    if isinstance(idx_or_text, int):
        p = d.paragraphs[idx_or_text]
        p.text = new_text
        return True
    for p in d.paragraphs:
        if idx_or_text in p.text:
            p.text = new_text
            return True
    return False

def clear_paras_containing(substrs):
    for p in d.paragraphs:
        for s in substrs:
            if s in p.text:
                p.text = ""
                break

def del_para(p):
    p._element.getparent().remove(p._element)

def set_cell(table_idx, r, c, text):
    t = d.tables[table_idx]
    while len(t.rows) <= r:
        t.add_row()
    row = t.rows[r]
    while len(row.cells) <= c:
        pass
    row.cells[c].text = text

def add_row(table_idx, values):
    t = d.tables[table_idx]
    row = t.add_row()
    for i, v in enumerate(values):
        if i < len(row.cells):
            row.cells[i].text = v
    return row

def insert_image_after_placeholder(placeholder_substr, image_path, width_inches=6.0, caption_after=None):
    for i, p in enumerate(d.paragraphs):
        if placeholder_substr in p.text and image_path in (p.text or ""):
            pass
    for p in d.paragraphs:
        if placeholder_substr == p.text.strip() or (placeholder_substr in p.text and len(p.text.strip()) < 80):
            p.text = ""
            run = p.add_run()
            run.add_picture(image_path, width=Inches(width_inches))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return True
    return False

def insert_image_at_para_with_text(match_text, image_path, width_inches=6.0, replace_text=True):
    for p in d.paragraphs:
        if match_text in p.text:
            if replace_text:
                p.text = ""
            run = p.add_run()
            run.add_picture(image_path, width=Inches(width_inches))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return True
    return False

# ============ 1. REMOVE GUIDELINE BLOCK (paras 0-20) ============
# Delete paragraphs containing guideline bullets + heading, keep document structure
guideline_markers = [
    "REPORT PREPARATION GUIDELINES",
    "These guidelines are provided",
    "Page setup: Use A4",
    "Font and headings:",
    "Paragraph format:",
    "Report structure:",
    "Template cleanup:",
    "Recommended length:",
    "Abstract and keywords:",
    "Present a precise problem statement",
    "Show the logical relational schema",
    "Demonstrate how the Java/Python",
    "ER diagram quality:",
    "Tables: Number tables",
    "Figures: Number figures",
    "Cross-references:",
    "Screenshots: Use clear",
    "Testing: Include meaningful",
    "References: Use IEEE",
    "Appendices: Place the complete",
    "Final submission: Check spelling",
]
for p in list(d.paragraphs):
    for m in guideline_markers:
        if m in p.text:
            del_para(p)
            break

# ============ 2. COVER PAGE ============
set_para("PROJECT TITLE", "Design and Implementation of a Database Management System\nfor E-Commerce Order, Return and Refund Management")
# Table 2: Program/Section/Faculty/Date
t2 = d.tables[2]
t2.rows[0].cells[1].text = "B.Tech \u2013 Artificial Intelligence & Machine Learning, School of Technology"
t2.rows[1].cells[1].text = "A"
t2.rows[2].cells[1].text = "DBMS Course Faculty, School of Technology, Woxsen University"
t2.rows[3].cells[1].text = "10 October 2026"
# row 4 academic year already 2026-27 (fix dash)
t2.rows[4].cells[1].text = "2026 \u2013 27"
# Table 3: student
t3 = d.tables[3]
t3.rows[1].cells[0].text = "25WU0102114"
t3.rows[1].cells[1].text = "Kakaraparthi Komal Prajvesh"
t3.rows[1].cells[2].text = ""
t3.rows[1].cells[3].text = ""
# Table 4: undertaking signature block
t4 = d.tables[4]
t4.rows[1].cells[3].text = "Name: Kakaraparthi Komal Prajvesh"
t4.rows[2].cells[3].text = "Roll Number: 25WU0102114"
t4.rows[3].cells[3].text = "Program: B.Tech AI & ML"

# Undertaking title placeholder
set_para("[Project Title]", "I hereby declare that the project entitled \u201cDesign and Implementation of a Database Management System for E-Commerce Order, Return and Refund Management\u201d, submitted as part of the Project-Based Learning activity for the Database Management Systems (DBMS) course, is my individual work carried out under the guidance of the course faculty.")
set_para("Wherever information, ideas, tools,",
         "I undertake that the project report, database design, SQL scripts, application source code, results, and supporting materials submitted by me have been prepared for academic purposes. Wherever information, ideas, tools, datasets, or other resources have been referred to, appropriate acknowledgement has been given. The MySQL 8 documentation, FastAPI documentation, and standard DBMS textbooks consulted are cited in the References [5]-[8].")

# ============ 3. ACRONYMS (Table 8) ============
acros = [
    ("3NF", "Third Normal Form"),
    ("ACID", "Atomicity, Consistency, Isolation, Durability"),
    ("API", "Application Programming Interface"),
    ("DBMS", "Database Management System"),
    ("DDL / DML", "Data Definition Language / Data Manipulation Language"),
    ("ENUM", "Enumerated (restricted status domain type in MySQL)"),
    ("ERD", "Entity-Relationship Diagram"),
    ("FK / PK", "Foreign Key / Primary Key"),
    ("JWT", "JSON Web Token"),
    ("RBAC", "Role-Based Access Control"),
    ("SPA", "Single Page Application"),
    ("SQL", "Structured Query Language"),
    ("PBL", "Project-Based Learning"),
    ("SRS", "Software Requirements Specification"),
]
t8 = d.tables[8]
# fill existing rows then add
for i, (a, e) in enumerate(acros):
    if i < len(t8.rows):
        t8.rows[i].cells[0].text = a
        t8.rows[i].cells[1].text = e
    else:
        add_row(8, [a, e])

# ============ 4. ABSTRACT + KEYWORDS ============
abstract = ("E-commerce post-purchase operations routinely fail when orders, returns, and refunds are tracked in disconnected spreadsheets, "
"causing untracked status changes, duplicate refunds, and missing audit evidence. The objective of this project is to design and implement a centralized ACID-compliant DBMS for the complete order-to-refund lifecycle. "
"The solution uses a 13-table Third Normal Form (3NF) MySQL 8 (InnoDB) schema covering customers, products, inventory, orders, payments, item-level returns, warehouse inspection, refunds, and an immutable audit log; five triggers enforce return-window eligibility, return-quantity limits, strict state transitions, and automatic auditing; and one stored procedure, sp_process_refund, executes atomic refund processing with duplicate-refund prevention and full rollback on error [5], [7]. "
"A Python 3.13 FastAPI backend with JWT role-based access (Customer, Support, Warehouse, Admin) and a responsive HTML/CSS/JS dashboard expose 17 REST endpoints for ordering, returning, approving, receiving, inspecting, refunding, auditing, and analytics [6], [8]. "
"The completed system demonstrates the full lifecycle Order \u2192 Requested \u2192 Approved \u2192 Received \u2192 Inspected \u2192 Refunded, passes 9/9 automated tests (constraints, procedure logic, integration, rollback, concurrency), sustains sub-millisecond indexed lookups at 1K scale, and verifies backup/restore integrity at 111,025 orders. "
"The application provides a single queryable source of truth for post-purchase operations with enforced business rules and traceable refunds.")
set_para("[Write the final abstract here.]", abstract)
set_para("Keywords: [Relational Database",
         "Keywords: E-Commerce Returns, Refund Management, MySQL, 3NF Normalization, Stored Procedures, Triggers, FastAPI")

# ============ 5. CHAPTER 1 ============
set_para("[Problem statement - approximately 1 to 2 paragraphs.]",
         "E-commerce after-sales handling is typically scattered across order spreadsheets, chat logs, and manual refund sheets. Support staff cannot see one dependable status for an order, warehouse staff record inspection results separately from the return approval, and finance has no systematic guard against duplicate refunds or late returns outside the product return window. Customers of stores such as the demo catalogue (headphones, keyboards, monitors, watches, docking stations) experience delayed or inconsistent refunds, while the business cannot answer analytical questions such as return rate per product or monthly refund totals. "
         "This project builds a single relational source of truth that links Customer \u2192 Orders \u2192 OrderItem \u2192 Payment \u2192 ReturnRequest \u2192 ReturnItem \u2192 Inspection \u2192 Refund, with every status change written to an immutable AuditLog. Database-enforced triggers block late returns, excess quantities, and invalid state jumps, and a transactional stored procedure guarantees that a refund is calculated from historical purchase prices exactly once or not at all.")

objectives = [
    "Design a 13-entity ER model and implement it as a normalized 3NF MySQL 8 (InnoDB) schema with PK/FK, CHECK, UNIQUE, ENUM, and ON DELETE RESTRICT integrity.",
    "Implement 5 triggers and 1 transactional stored procedure (sp_process_refund) to automate return eligibility, state-machine enforcement, audit logging, and atomic refunds.",
    "Develop a FastAPI + JWT RBAC application with 17 REST endpoints and a dashboard that performs all CRUD through parameterized SQL on the designed database.",
    "Validate the system with 9 automated tests (100% pass) covering constraints, procedure logic, full lifecycle integration, rollback, and concurrent duplicate-refund handling.",
    "Generate operational and analytical reports (order history, pending returns, refund status, return rates, monthly refunds, top reasons) and verify performance up to 100K orders with backup/restore checks.",
]
# objectives are List Bullet paras [Objective 1..4]; replace first 4, add 5th
count = 0
for p in d.paragraphs:
    if p.text.strip() in ("[Objective 1]", "[Objective 2]", "[Objective 3]", "[Objective 4]"):
        p.text = objectives[count]
        count += 1
# add 5th objective after the 4th bullet: find para with objectives[3]
for i, p in enumerate(d.paragraphs):
    if p.text == objectives[3]:
        new_p = deepcopy(p._element)
        # insert after
        p._element.addnext(new_p)
        d.paragraphs[i+1].text if False else None
        break
# simpler: set via python-docx insert not trivial; instead append to same para list by adding new paragraph in document body near? We'll just add bullet at end of doc and move on:
# Workaround: find the element and insert new <w:p>
from docx.oxml import OxmlElement
# skip complexity: append objective 5 to 4th bullet para is bad. Instead create new paragraph after objectives[3] using insert_paragraph_before on next heading
for p in d.paragraphs:
    if "Scope and Assumptions" in p.text and "1.3" in p.text:
        # insert before this heading
        new = OxmlElement('w:p')
        p._element.addprevious(new)
        break
# Fallback simpler: just extend 4th objective text to include 5th? No — instead handle by editing: put objectives[4] appended as separate sentence in 4th bullet is wrong.
# Let's do it properly with python-docx: insert paragraph
inserted = False
for p in list(d.paragraphs):
    if p.text == objectives[3] and not inserted:
        # create new paragraph element after p
        from docx.oxml.ns import qn
        new_p_el = deepcopy(p._element)
        # set text of copy
        for t_el in new_p_el.findall('.//w:t', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}):
            pass
        p._element.addnext(new_p_el)
        inserted = True
        break
# set text of the duplicated paragraph (it is now paragraphs list changed; find duplicate)
found_first = False
for p in d.paragraphs:
    if p.text == objectives[3]:
        if not found_first:
            found_first = True
        else:
            p.text = objectives[4]
            break

# Table 10 scope (Table index 10)
t10 = d.tables[10]
t10.rows[0].cells[1].text = "Order placement with mock payment; item-level return requests; support approve/reject; warehouse receive + inspection (GOOD/DAMAGED/DEFECTIVE); atomic refund via sp_process_refund; inventory restock on GOOD; immutable audit log; 6 operational/analytical reports; 4-role RBAC demo."
t10.rows[1].cells[1].text = "Real bank/card network integration (mock gateway only); multi-region replication; image-based auto-inspection; SMS/email webhooks; mobile app; distributed sharding."
t10.rows[2].cells[1].text = "Single MySQL 8 InnoDB instance on localhost/Docker; all prices in one currency; return window counted from Orders.order_date per product return_window_days; seeded demo users share password password123; tested to 100K orders locally."

# Table 11 solution strategy (index 11)
t11 = d.tables[11]
t11.rows[0].cells[1].text = "Customer (places orders, requests returns); Support Agent (approves/rejects, triggers refunds); Warehouse Operator (receives, inspects); Admin (full access + audit log)."
t11.rows[1].cells[1].text = "Auth/JWT login; Customer portal (products, orders, returns); Operations dashboard (support + warehouse queues); Refund engine (sp_process_refund); Analytics (return rates, monthly refunds, top reasons); Audit viewer."
t11.rows[2].cells[1].text = "Customer orders (Delivered) \u2192 requests return (Requested) \u2192 Support approves (Approved) \u2192 Warehouse receives (Received) \u2192 inspects (Inspected) \u2192 Support/Admin refunds (Refunded); Rejected is terminal from any non-terminal state; triggers block any skip."
t11.rows[3].cells[1].text = "Return only within product return window; return qty \u2264 purchased \u2212 already returned; strict state machine Requested\u2192Approved/Rejected\u2192Received\u2192Inspected\u2192Refunded/Rejected; one processed refund per return; GOOD inspection restocks inventory; every change audited."
t11.rows[4].cells[1].text = "Order history with returns; pending returns queue; refund status by order; return rate per product; monthly refund totals; top return reasons; Admin audit log (last 100 events)."

# Table 12 FR (index 12)
frs = [
    ("FR-01", "Customer places order with items; system prices from Product, creates Orders + OrderItem + Payment (mock gateway reference).", "High"),
    ("FR-02", "Customer submits item-level return request (reason + quantities) within return window.", "High"),
    ("FR-03", "Support approves or rejects a Requested return (state trigger enforces).", "High"),
    ("FR-04", "Warehouse marks Approved return as Received on physical arrival.", "High"),
    ("FR-05", "Warehouse records Inspection (GOOD/DAMAGED/DEFECTIVE + notes); return becomes Inspected.", "High"),
    ("FR-06", "Support/Admin processes refund exactly once via sp_process_refund (amount = \u03a3 qty \u00d7 unit_price_at_purchase); duplicate blocked, rollback on error.", "High"),
    ("FR-07", "Admin views immutable AuditLog of all ReturnRequest/Refund status changes with actor and timestamp.", "Medium"),
    ("FR-08", "All roles view analytics: order history, pending returns, refund status, return rates, monthly refunds, top reasons (RBAC-filtered endpoints).", "Medium"),
]
t12 = d.tables[12]
for i, (fid, req, pri) in enumerate(frs):
    if i < len(t12.rows) - 1:
        t12.rows[i+1].cells[0].text = fid
        t12.rows[i+1].cells[1].text = req
        t12.rows[i+1].cells[2].text = pri
    else:
        add_row(12, [fid, req, pri])

# Table 13 entity summary (index 13)
ents = [
    ("Role", "role_id (PK), role_name (UK)", "RBAC roles: Customer, Support, Warehouse, Admin."),
    ("Customer", "customer_id (PK), email (UK)", "Buyer profile, separate from login credentials."),
    ("AppUser", "user_id (PK), username (UK); FK role_id, customer_id", "Login credentials + role assignment; bcrypt hash."),
    ("Product", "product_id (PK), current_price, return_window_days", "Catalogue master data; window 7/14/30 days."),
    ("Inventory", "inventory_id (PK), serial_or_sku (UK); FK product_id", "Stock units; status AVAILABLE/RESERVED/SOLD/RETURNED/DAMAGED."),
    ("Orders", "order_id (PK); FK customer_id", "Order header; total_amount; status Pending..Cancelled."),
    ("OrderItem", "order_item_id (PK); FK order_id, product_id", "Line items; preserves unit_price_at_purchase (history)."),
    ("Payment", "payment_id (PK), provider_reference (UK); FK order_id", "Mock gateway record; no raw card data stored."),
    ("ReturnRequest", "return_id (PK); FK order_id, customer_id", "Return header; 6-state lifecycle; see Fig. lifecycle."),
    ("ReturnItem", "return_item_id (PK); FK return_id, order_item_id", "Partial returns; qty > 0 enforced by trigger."),
    ("Inspection", "inspection_id (PK), return_id (UK, 1:1); FK return_id", "Warehouse verdict GOOD/DAMAGED/DEFECTIVE + inspector."),
    ("Refund", "refund_id (PK); FK return_id, payment_id", "Financial credit; 1 return : N refunds allowed, 1 processed enforced."),
    ("AuditLog", "log_id (PK); entity_type + entity_id", "Immutable trail (old/new status, actor, note, timestamp)."),
    ("Customer 1:N Orders; Orders 1:N OrderItem", "Cardinality 1:N, total on N side", "A customer places many orders; an order has 1..N items."),
    ("Orders 1:N Payment; ReturnRequest 1:1 Inspection", "1:N and 1:1 (UNIQUE)", "Payment per order; exactly one inspection per return."),
    ("ReturnRequest 1:N Refund; OrderItem 1:N ReturnItem", "1:N", "Refunds itemised from purchase history; partial returns supported."),
]
t13 = d.tables[13]
# header row assumed; clear existing data rows
for i in range(1, len(t13.rows)):
    for c in t13.rows[i].cells:
        c.text = ""
for i, (a, b, c_) in enumerate(ents):
    if i + 1 < len(t13.rows):
        t13.rows[i+1].cells[0].text = a
        t13.rows[i+1].cells[1].text = b
        t13.rows[i+1].cells[2].text = c_
    else:
        add_row(13, [a, b, c_])

# ER diagram image
insert_image_at_para_with_text("[Insert approved ER diagram here]", ER_IMG, width_inches=6.2)

print("Chapter 1 done")
d.save(OUT)
print("saved intermediate")
